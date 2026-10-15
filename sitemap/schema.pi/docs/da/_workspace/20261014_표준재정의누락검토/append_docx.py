# 신규 DOCX 끝에 "보완 절"을 덧붙인다 — 저장소 루트에서 python append_docx.py <spec.json> <백업 디렉토리> (spec 의 file 은 저장소 루트 기준 상대경로)
# spec: [{"file": "<docx 경로>", "sections": [{"heading": "...", "level": 1|2|3,
#          "paras": ["..."], "table": {"header": ["..."], "rows": [["..."]]}}]}]
# 기존 본문은 건드리지 않고 마지막에만 추가(출처 병기는 spec 문장에 포함). 원본은 백업 디렉토리에 1회 보관(이미 있으면 덮어쓰지 않음)
import json, sys, shutil
from pathlib import Path
import docx
from docx.shared import Pt

spec = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
bakdir = Path(sys.argv[2]); bakdir.mkdir(parents=True, exist_ok=True)
for item in spec:
    p = Path(item['file'])
    bak = bakdir / p.name
    if not bak.exists():
        shutil.copy2(p, bak)
    d = docx.Document(p)
    # 기존 제목 스타일 이름 확인(한글 템플릿은 Heading 1 이 없을 수 있음)
    styles = {s.name for s in d.styles}
    def hstyle(level):
        name = f'Heading {level}'
        return name if name in styles else None
    for sec in item['sections']:
        st = hstyle(sec.get('level', 1))
        if st:
            d.add_paragraph(sec['heading'], style=st)
        else:
            para = d.add_paragraph()
            r = para.add_run(sec['heading']); r.bold = True; r.font.size = Pt(13)
        for t in sec.get('paras', []):
            d.add_paragraph(t)
        tb = sec.get('table')
        if tb:
            table = d.add_table(rows=1, cols=len(tb['header']))
            table.style = 'Table Grid' if 'Table Grid' in styles else table.style
            for i, h in enumerate(tb['header']):
                table.rows[0].cells[i].text = h
            for row in tb['rows']:
                cells = table.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
            d.add_paragraph()
    d.save(p)
    print('appended', p.name, 'sections', len(item['sections']))
