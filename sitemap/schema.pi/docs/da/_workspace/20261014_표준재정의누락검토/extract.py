# DOCX·PPTX·XLS → 평문 .md 변환 — python extract.py <입력 디렉토리> <출력 디렉토리>
# 표는 행마다 "| 셀 | 셀 |", 제목 스타일은 '#' 로 바꿔 팀원이 절 단위로 읽을 수 있게 한다
import sys, os, re
from pathlib import Path

def docx_to_md(p):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(p)
    out = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            para = Paragraph(child, d)
            t = para.text.strip()
            if not t:
                continue
            st = (para.style.name or '').lower()
            m = re.match(r'heading (\d)', st)
            if m:
                out.append('#' * min(int(m.group(1)), 6) + ' ' + t)
            elif st.startswith('title'):
                out.append('# ' + t)
            else:
                out.append(t)
        elif tag == 'tbl':
            tb = Table(child, d)
            for row in tb.rows:
                cells = []
                prev = None
                for c in row.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(' '.join(c.text.split()))
                out.append('| ' + ' | '.join(cells) + ' |')
            out.append('')
    return '\n'.join(out)

def pptx_to_md(p):
    from pptx import Presentation
    pr = Presentation(p)
    out = []
    for i, s in enumerate(pr.slides, 1):
        out.append(f'## 슬라이드 {i}')
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = ''.join(r.text for r in para.runs).strip()
                    if t:
                        out.append(('  ' * para.level) + '- ' + t)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    out.append('| ' + ' | '.join(' '.join(c.text.split()) for c in row.cells) + ' |')
        out.append('')
    return '\n'.join(out)

def xls_to_md(p):
    import xlrd
    wb = xlrd.open_workbook(p)
    out = []
    for sh in wb.sheets():
        out.append(f'## 시트 {sh.name} ({sh.nrows}행)')
        for r in range(sh.nrows):
            vals = [' '.join(str(sh.cell_value(r, c)).split()) for c in range(sh.ncols)]
            if any(vals):
                out.append('| ' + ' | '.join(vals) + ' |')
        out.append('')
    return '\n'.join(out)

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
dst.mkdir(parents=True, exist_ok=True)
for f in sorted(src.iterdir()):
    if f.is_dir() or f.name.startswith('~$'):
        continue
    ext = f.suffix.lower()
    try:
        if ext == '.docx': md = docx_to_md(f)
        elif ext == '.pptx': md = pptx_to_md(f)
        elif ext == '.xls': md = xls_to_md(f)
        else: continue
    except Exception as e:
        print('FAIL', f.name, e); continue
    o = dst / (f.stem + '.md')
    o.write_text(md, encoding='utf-8')
    print(f'{f.name} -> {o.name} {len(md)} chars {md.count(chr(10))} lines')
