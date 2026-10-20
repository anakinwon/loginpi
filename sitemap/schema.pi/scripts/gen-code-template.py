# 표준코드(공통코드) 일괄등록 엑셀 템플릿 생성기(2026-10-20) — public/templates/std-code-bulk-template.xlsx 를 만든다(정본 = 이 파일, xlsx 직접 수정 금지)
# 사용 : python scripts/gen-code-template.py public/templates
# 헤더 문자열은 src/lib/bulk.ts HEADERS.codes 와 같아야 한다(일괄등록 화면이 헤더 이름으로 열을 찾음). 화면은 CSV(UTF-8)만 읽으므로 작성 뒤 'CSV UTF-8' 로 저장해 올린다
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else 'public/templates'
ROWS = 1000  # 입력 검사(사용여부·정렬순서)를 거는 행 수
C = {'head': '1F3A68', 'req': 'C0392B', 'band': 'DCE6F2', 'ex': 'F3F4F6', 'line': 'B8C6E0'}
thin = Side(style='thin', color=C['line'])
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# (헤더, 필수, 너비, 설명) — 순서 = 화면 안내 순서
COLS = [
    ('코드명', True, 22, '코드기본 = 확정 표준용어(코드형 도메인). 이름은 ~구분코드·~종류코드·~유형코드 3종만'),
    ('코드분류', False, 14, '코드를 묶는 분류(예 공통·결제). 코드기본을 새로 만들 때 그 코드의 첫 행 값만 씀'),
    ('코드값', True, 12, '코드상세 값. 용어 도메인 길이 이내, 같은 코드 안 중복 불가'),
    ('코드값명', True, 22, '코드값의 이름. 같은 코드 안 중복 불가(상호 배타)'),
    ('정렬순서', False, 10, '0~99999 정수, 비우면 0'),
    ('사용여부', False, 10, 'Y 또는 N, 비우면 Y'),
    ('코드값설명', False, 40, '코드값 설명(4000자 이내)'),
]
EXAMPLES = [
    ('남녀구분코드', '공통', 'M', '남자', 1, 'Y', '구분코드 — 값이 고정적'),
    ('남녀구분코드', '공통', 'F', '여자', 2, 'Y', ''),
    ('지불수단종류코드', '결제', 'CARD', '신용카드', 1, 'Y', '종류코드 — 계속 늘어날 수 있음'),
    ('지불수단종류코드', '결제', 'BANK', '계좌이체', 2, 'Y', ''),
    ('지불수단종류코드', '결제', 'PI', '파이', 3, 'Y', ''),
    ('거래유형코드', '결제', 'BUY', '구매', 1, 'Y', '유형코드 — 성질이 유사한 것끼리 묶음'),
    ('거래유형코드', '결제', 'RFND', '환불', 2, 'Y', ''),
]


def header(ws, row=1):
    for i, (nm, req, w, desc) in enumerate(COLS, 1):
        c = ws.cell(row=row, column=i, value=nm)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=C['req'] if req else C['head'])
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = box
        c.comment = Comment(('[필수] ' if req else '') + desc, 'schema.pi')
        ws.column_dimensions[c.column_letter].width = w


wb = Workbook()
# 1) 입력 시트 — 첫 시트를 CSV 로 저장해 올림(헤더 1행 + 데이터)
ws = wb.active
ws.title = '표준코드'
header(ws)
ws.freeze_panes = 'A2'
yn = DataValidation(type='list', formula1='"Y,N"', allow_blank=True, error='Y 또는 N', errorTitle='사용여부')
seq = DataValidation(type='whole', operator='between', formula1='0', formula2='99999', allow_blank=True, error='0~99999 정수', errorTitle='정렬순서')
ws.add_data_validation(yn)
ws.add_data_validation(seq)
yn.add(f'F2:F{ROWS + 1}')
seq.add(f'E2:E{ROWS + 1}')
for col in ('A', 'C'):
    for r in range(2, ROWS + 2):
        ws[f'{col}{r}'].number_format = '@'  # 코드값 01 같은 앞자리 0 보존(텍스트)

# 2) 작성 안내 시트
g = wb.create_sheet('작성안내')
g.column_dimensions['A'].width = 18
g.column_dimensions['B'].width = 100
lines = [
    ('표준코드(공통코드) 일괄등록 템플릿', None),
    ('공통코드 이름', '~구분코드 · ~종류코드 · ~유형코드 3종만 허용(그 외 코드는 공통코드가 아니므로 별도 기본테이블로 관리)'),
    ('구분코드', '코드 명이 더 이상 늘어나지 않고 고정적일 때 사용. 예) 남녀구분코드'),
    ('종류코드', '고정적이지 않고 지속적으로 늘어날 수 있을 때 사용. 예) 서비스종류코드, 지불수단종류코드'),
    ('유형코드', '성질이나 특징이 유사한 것끼리 묶을 때 사용. 예) 지불수단유형코드, 거래유형코드'),
    ('구조', '코드기본 : 코드상세 = 1 : N — 한 행 = 코드상세 하나. 코드기본이 없으면 반영할 때 함께 등록'),
    ('사전 조건', '코드명은 표준용어로 먼저 등록·승인(확정)되어 있어야 하고, 도메인이 코드형이어야 함'),
    ('설계 원칙', '코드상세는 성격 동일 · 상호 배타(코드값·코드값명 중복 불가) · 전체 합집합 · 상세화 수준 유사'),
    ('저장 방법', "'표준코드' 시트에 작성 → 파일 › 다른 이름으로 저장 › 'CSV UTF-8(쉼표로 분리)(*.csv)' → 표준코드 › 일괄등록 화면에서 파일 선택 → 검증 → 반영"),
    ('주의', '헤더(1행) 이름을 바꾸지 마십시오. 코드값 칸은 텍스트 형식(앞자리 0 보존). 이미 있는 코드값은 건너뜀, 승인 없이 바로 반영'),
]
for i, (k, v) in enumerate(lines, 1):
    a = g.cell(row=i, column=1, value=k)
    a.font = Font(bold=True, size=14 if v is None else 11, color=C['head'])
    if v is not None:
        a.fill = PatternFill('solid', fgColor=C['band'])
        a.border = box
        b = g.cell(row=i, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical='center')
        b.border = box
r0 = len(lines) + 2
g.cell(row=r0, column=1, value='칸 설명').font = Font(bold=True, size=12, color=C['head'])
for i, (nm, req, _, desc) in enumerate(COLS, r0 + 1):
    a = g.cell(row=i, column=1, value=nm + (' (필수)' if req else ''))
    a.font = Font(bold=True, color=C['req'] if req else C['head'])
    a.border = box
    b = g.cell(row=i, column=2, value=desc)
    b.border = box

# 3) 작성 예시 시트(그대로 올리면 표준용어가 없을 때 '용어 없음' 오류 — 형식 참고용)
e = wb.create_sheet('작성예시')
header(e)
e.freeze_panes = 'A2'
for r, row in enumerate(EXAMPLES, 2):
    for c, v in enumerate(row, 1):
        cell = e.cell(row=r, column=c, value=v)
        cell.fill = PatternFill('solid', fgColor=C['ex'])
        cell.border = box
        if c in (1, 3):
            cell.number_format = '@'

wb.save(f'{OUT}/std-code-bulk-template.xlsx')
print('ok', f'{OUT}/std-code-bulk-template.xlsx')
