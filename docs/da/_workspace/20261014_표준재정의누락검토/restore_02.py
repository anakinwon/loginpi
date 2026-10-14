# 신규 02 지침서 끝에 구 데이터표준화지침서 3~6장을 복원(안 A, 10_standards_gap §3) — python restore_02.py
# 원문은 그대로 옮기되 '000㈜'→'기관' 치환, 그림은 이미지 파트를 함께 복사. 조정 사항은 본문을 고치지 않고 끝에 표로 둔다
import copy, io, re, shutil
from pathlib import Path
import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]  # docs/da
OLD = ROOT / 'references/backup/데이터표준화지침서.docx'
NEW = ROOT / 'references/02. 데이터 표준화 지침서.docx'
BAK = Path(__file__).resolve().parent / 'bak' / NEW.name
if not BAK.exists():
    shutil.copy2(NEW, BAK)

old = docx.Document(OLD)
new = docx.Document(NEW)
kids = list(old.element.body.iterchildren())
# 3장 제목(Heading 1 "3. 표준 단어") 부터 sectPr 직전까지
start = next(i for i, c in enumerate(kids) if c.tag == qn('w:p') and Paragraph(c, old).text.strip().startswith('3. 표준 단어') and 'Heading' in (Paragraph(c, old).style.name or ''))
end = len(kids) - 1 if kids[-1].tag == qn('w:sectPr') else len(kids)
body = new.element.body
sect = body.find(qn('w:sectPr'))

def add_before_sect(el):
    if sect is not None:
        sect.addprevious(el)
    else:
        body.append(el)

h = new.add_paragraph('구 데이터표준화지침서 3~6장 복원 [보완 2026-10-14]', style='Heading 1')
p = new.add_paragraph('이 지침서의 「적용 원칙과 조정 사항」이 "각 장의 본문에도 반영했다"고 적고 있으나 3~6장 본문이 빠져 있고, 「정의서」·「모델링 가이드」·「실습 사례집」이 3.3·4.3·5.4·6.5·6.7 등 이 장들을 44회 인용한다. 구 데이터표준화지침서 3~6장(= 10·20·30·40 지침서와 동일)을 원문 그대로 되살리고, 필수 표준(STD-01~08)·정본과 다른 곳은 본문을 고치지 않고 끝의 "복원본 조정 사항" 표로 정한다. 읽을 때는 조정 표가 본문보다 우선한다. (출처: 구 데이터표준화지침서 3~6장, 10_standards_gap §3 안 A)')
# add_paragraph 는 sectPr 앞에 넣어 주므로 그대로 둔다

img_map = {}
n_img = 0
for c in kids[start:end]:
    el = copy.deepcopy(c)
    for t in el.iter(qn('w:t')):
        if t.text and '000㈜' in t.text:
            t.text = t.text.replace('000㈜', '기관')
    for blip in el.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        if not rid:
            continue
        if rid not in img_map:
            part = old.part.related_parts[rid]
            new_rid, _ = new.part.get_or_add_image(io.BytesIO(part.blob))
            img_map[rid] = new_rid
            n_img += 1
        blip.set(qn('r:embed'), img_map[rid])
    add_before_sect(el)

new.add_paragraph('복원본 조정 사항 [보완 2026-10-14]', style='Heading 2')
new.add_paragraph('위 복원 본문 가운데 아래 항목은 조정 결과대로 읽는다. 「적용 원칙과 조정 사항」 조정표 7행도 그대로 적용된다. (출처: 10_standards_gap §3-3·§2-3)')
rows = [
    ('데이터타입 분류 CH', 'CH 삭제, 문자형은 VC 일원화. IV(Interval) 추가', '조정표 1행, STD-01'),
    ('날짜VC8', '삭제. 일자·일시는 DATE(날짜DT)', '조정표 2행, STD-02'),
    ('부분 날짜 예외', '연도·연월·월일·시분초·연도분기·연도반기 외에 월(VC2)·일(VC2)도 예외에 포함(4.3 예시표 결산월·결제일·기준일)', '조정표 3행, D40'),
    ('타임스탬프 TS', '로그성 속성에 한해 DA 승인 후 사용', '조정표 4행'),
    ('여부 값', '모름·해당없음 NULL(5.4) vs 0·9(4.3) 병존 → 0·9, NOT NULL DEFAULT N (PostgreSQL 프로파일은 CHAR(1)+CHECK)', '조정표 5행, STD-06, T66'),
    ('주제영역 약어', '4자리 → 3자리(시스템 1 + 1레벨 2)', '조정표 6행'),
    ('영문 약어 길이', '신규 등록 약어는 자음 위주 2~3자를 원칙으로 하고 4자는 관용 약어에 한해 허용. 기존 4자 이상 약어는 유지', '3.3, 데이터표준규칙 §1-1 v2.3, SC03'),
    ('도메인 약어 길이', '2자 원칙, 일시(DTM)처럼 2자로 구분되지 않는 관용형은 3자까지 허용. 이 프로젝트의 도메인 약어는 데이터표준규칙 §1-2 표를 따른다', '3.3, SC04'),
    ("'순서(ORDER) OD' 예시", "삭제. ORDER는 주문 ORD로 확정(데이터표준규칙 §1-1 v2.3). 순서 도메인은 폐지하고 순번 seq로 통합", '3.3, W56·W68·SC02'),
    ('1개월 약어', 'MM1(복합어 절)과 MT1(약어 절) 병존 → MM1로 통일(실습 사례집과 일치)', '3.3, W71'),
    ("사람 접미사 사례표", "'신청자→신청인'과 '신청인→신청자' 동시 수록 → 신청인으로 통일. '조사원→제조사'는 오기이므로 무시", '3.3, W22'),
    ('금액 타입', 'NB24.3(4.3 본문)과 NB21.3(4.2 예시) 병존 → NB21.3으로 통일', '4.2·4.3, D33·SC09'),
    ('HHMMSS 도메인', "시분초(4.3)와 시각(5.4) 병존 → 시분초로 통일. 5.4 '전문전송시각' 예시는 '전문전송시분초'로 읽음", '4.3·5.4, SC10'),
    ('두 식별번호 통합 용어', "'주민사업자등록번호'(4.3)와 '주민사업자등록통합번호'(5.4) 병존 → 통합번호로 통일", '4.3·5.4, SC11·T56'),
    ('코드값명 문자 규칙', '코드값명에는 표준단어가 아닌 말도 쓸 수 있으나 가능하면 표준단어를 쓴다. 한글·영문·숫자와 특수문자 / + & ( ) _ % 허용 (40.표준코드_지침서에만 있던 1문장 추가)', '6.7, K44'),
    ("'000㈜' 표기", "'기관'으로 치환해 복원", '전 장'),
    ('그림', '구 원본의 그림 %d개를 이미지 파트째 복사. 깨진 그림이 있으면 구 지침서 원본(references/backup)을 본다' % n_img, '전 장'),
]
tb = new.add_table(rows=1, cols=3)
try:
    tb.style = 'Table Grid'
except Exception:
    pass
for i, hd in enumerate(['항목', '조정 결과', '근거']):
    tb.rows[0].cells[i].text = hd
for r in rows:
    cells = tb.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = v
new.save(NEW)
print('restored children', end - start, 'images', n_img)
